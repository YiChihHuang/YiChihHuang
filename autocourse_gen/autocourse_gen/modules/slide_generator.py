"""
Slide Generator Module
======================

Generates PowerPoint slides from the course curriculum.
"""

from pathlib import Path
from typing import Optional, Dict, Any, List

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from rich.console import Console
from rich.panel import Panel
from rich.progress import track

from ..utils.resumable import check_and_resume

console = Console()


def hex_to_rgb(hex_color: str) -> RgbColor:
    """Convert hex color to RgbColor."""
    hex_color = hex_color.lstrip("#")
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    return RgbColor(r, g, b)


class SlideGeneratorModule:
    """
    Slide Generator Module for creating PowerPoint presentations.

    Creates clean, minimal slides optimized for video recording
    with space for mouse cursor movement.
    """

    OUTPUT_FILE = "3_slides.pptx"

    def __init__(self, workspace: Path, config: dict):
        """
        Initialize the Slide Generator Module.

        Args:
            workspace: Workspace directory for output
            config: Slide configuration settings
        """
        self.workspace = workspace
        self.config = config
        self.output_path = workspace / self.OUTPUT_FILE

        # Slide dimensions
        self.width = Inches(config.get("width_inches", 13.333))
        self.height = Inches(config.get("height_inches", 7.5))

        # Colors
        self.primary_color = hex_to_rgb(config.get("primary_color", "#2B579A"))
        self.secondary_color = hex_to_rgb(config.get("secondary_color", "#5B9BD5"))
        self.accent_color = hex_to_rgb(config.get("accent_color", "#ED7D31"))

        # Fonts
        self.title_font = config.get("title_font", "Microsoft JhengHei")
        self.body_font = config.get("body_font", "Microsoft JhengHei")

    def run(
        self,
        syllabus: Dict[str, Any],
        skip_if_exists: bool = True,
    ) -> Path:
        """
        Generate PowerPoint slides from syllabus.

        Args:
            syllabus: Course syllabus dictionary
            skip_if_exists: Whether to skip if output file exists

        Returns:
            Path to generated PPTX file
        """
        console.print(Panel(
            "[bold blue]Step 3: Slide Generation[/bold blue]\n"
            "Creating PowerPoint presentation...",
            title="🎨 Slide Generator",
        ))

        # Check for existing output
        if skip_if_exists:
            should_skip, _ = check_and_resume(
                self.output_path,
                "Slide Generation",
                ask_user=True,
            )
            if should_skip:
                return self.output_path

        # Create presentation
        prs = Presentation()
        prs.slide_width = self.width
        prs.slide_height = self.height

        # Generate slides
        self._create_title_slide(prs, syllabus)
        self._create_overview_slide(prs, syllabus)

        sections = syllabus.get("sections", [])
        for section in track(sections, description="Generating slides..."):
            self._create_section_slides(prs, section)

        self._create_summary_slide(prs, syllabus)

        # Save presentation
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(self.output_path))

        # Display result
        console.print("\n[green]✓ Slides generated successfully![/green]")
        console.print(f"  Output saved to: [cyan]{self.output_path}[/cyan]")
        console.print(f"  Total slides: {len(prs.slides)}")
        console.print("\n[dim]Tip: You can open the PPTX file in PowerPoint to make manual adjustments.[/dim]")

        return self.output_path

    def _create_title_slide(self, prs: Presentation, syllabus: Dict[str, Any]):
        """Create the title slide."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0, 0, self.width, self.height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self.primary_color
        background.line.fill.background()

        # Title
        title = syllabus.get("course_title", "課程")
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5),
            self.width - Inches(1), Inches(2)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = RgbColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        # Subtitle
        description = syllabus.get("course_description", "")
        if description:
            subtitle_box = slide.shapes.add_textbox(
                Inches(1), Inches(4.5),
                self.width - Inches(2), Inches(1)
            )
            tf = subtitle_box.text_frame
            p = tf.paragraphs[0]
            p.text = description
            p.font.size = Pt(24)
            p.font.color.rgb = RgbColor(220, 220, 220)
            p.alignment = PP_ALIGN.CENTER

    def _create_overview_slide(self, prs: Presentation, syllabus: Dict[str, Any]):
        """Create the course overview slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        self._add_slide_title(slide, "課程大綱")

        # Content - list of sections
        sections = syllabus.get("sections", [])
        content_box = slide.shapes.add_textbox(
            Inches(1), Inches(1.8),
            self.width - Inches(2), Inches(5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        for i, section in enumerate(sections):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"{section.get('section_id', i+1)}. {section.get('title', 'Section')}"
            p.font.size = Pt(28)
            p.font.color.rgb = self.primary_color
            p.space_before = Pt(16)
            p.space_after = Pt(8)

            # Add brief description
            if section.get("description"):
                p2 = tf.add_paragraph()
                p2.text = f"    {section['description'][:60]}..."
                p2.font.size = Pt(18)
                p2.font.color.rgb = RgbColor(100, 100, 100)

    def _create_section_slides(self, prs: Presentation, section: Dict[str, Any]):
        """Create slides for a section."""
        # Section title slide
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Section header background
        header_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0, 0, self.width, Inches(2)
        )
        header_bg.fill.solid()
        header_bg.fill.fore_color.rgb = self.secondary_color
        header_bg.line.fill.background()

        # Section number and title
        section_id = section.get("section_id", "?")
        title = section.get("title", "Section")

        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5),
            self.width - Inches(1), Inches(1.2)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Chapter {section_id}"
        p.font.size = Pt(20)
        p.font.color.rgb = RgbColor(255, 255, 255)

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(40)
        p2.font.bold = True
        p2.font.color.rgb = RgbColor(255, 255, 255)

        # Key concepts
        key_concepts = section.get("key_concepts", [])
        if key_concepts:
            concepts_box = slide.shapes.add_textbox(
                Inches(1), Inches(2.5),
                self.width - Inches(2), Inches(4)
            )
            tf = concepts_box.text_frame
            tf.word_wrap = True

            p = tf.paragraphs[0]
            p.text = "本章重點"
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = self.primary_color

            for concept in key_concepts:
                p = tf.add_paragraph()
                p.text = f"• {concept}"
                p.font.size = Pt(22)
                p.space_before = Pt(12)

        # Metaphor slide (if present)
        if section.get("metaphor"):
            self._create_metaphor_slide(prs, section)

        # Talking points slides
        talking_points = section.get("talking_points", [])
        for point in talking_points:
            self._create_content_slide(prs, section, point)

    def _create_metaphor_slide(self, prs: Presentation, section: Dict[str, Any]):
        """Create a slide for the section metaphor."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        self._add_slide_title(slide, "💡 用比喻來理解")

        # Metaphor content
        metaphor = section.get("metaphor", "")
        content_box = slide.shapes.add_textbox(
            Inches(1), Inches(2),
            self.width - Inches(2), Inches(4.5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = metaphor
        p.font.size = Pt(28)
        p.font.color.rgb = RgbColor(60, 60, 60)
        p.alignment = PP_ALIGN.CENTER

    def _create_content_slide(
        self,
        prs: Presentation,
        section: Dict[str, Any],
        point: Dict[str, Any]
    ):
        """Create a content slide for a talking point."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        topic = point.get("topic", "Topic")
        self._add_slide_title(slide, topic)

        # Content
        content = point.get("content", "")
        content_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.8),
            self.width - Inches(1.6), Inches(5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        # Split content into bullet points if it's long
        if len(content) > 100:
            # Try to split by sentences or newlines
            sentences = content.replace("。", "。\n").split("\n")
            sentences = [s.strip() for s in sentences if s.strip()]

            for i, sentence in enumerate(sentences[:6]):  # Limit to 6 bullets
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = f"• {sentence}" if not sentence.startswith("•") else sentence
                p.font.size = Pt(22)
                p.font.color.rgb = RgbColor(50, 50, 50)
                p.space_before = Pt(12)
        else:
            p = tf.paragraphs[0]
            p.text = content
            p.font.size = Pt(24)
            p.font.color.rgb = RgbColor(50, 50, 50)

        # Notes (if present)
        if point.get("notes"):
            notes_box = slide.shapes.add_textbox(
                Inches(0.8), Inches(6.5),
                self.width - Inches(1.6), Inches(0.8)
            )
            tf = notes_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"📝 {point['notes']}"
            p.font.size = Pt(16)
            p.font.italic = True
            p.font.color.rgb = RgbColor(120, 120, 120)

    def _create_summary_slide(self, prs: Presentation, syllabus: Dict[str, Any]):
        """Create the summary slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0, 0, self.width, self.height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self.primary_color
        background.line.fill.background()

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1),
            self.width - Inches(1), Inches(1)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = "課程總結"
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RgbColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        # Summary content
        summary = syllabus.get("summary", "感謝您的學習！")
        summary_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.5),
            self.width - Inches(2), Inches(2)
        )
        tf = summary_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = summary
        p.font.size = Pt(24)
        p.font.color.rgb = RgbColor(230, 230, 230)
        p.alignment = PP_ALIGN.CENTER

        # Next steps
        next_steps = syllabus.get("next_steps", [])
        if next_steps:
            steps_box = slide.shapes.add_textbox(
                Inches(1), Inches(4.5),
                self.width - Inches(2), Inches(2.5)
            )
            tf = steps_box.text_frame
            tf.word_wrap = True

            p = tf.paragraphs[0]
            p.text = "📈 延伸學習"
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = RgbColor(255, 255, 255)

            for step in next_steps[:3]:
                p = tf.add_paragraph()
                p.text = f"• {step}"
                p.font.size = Pt(18)
                p.font.color.rgb = RgbColor(200, 200, 200)
                p.space_before = Pt(8)

    def _add_slide_title(self, slide, title: str):
        """Add a title to a slide."""
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3),
            self.width - Inches(1), Inches(1)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = self.primary_color

        # Underline
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5), Inches(1.3),
            Inches(2), Pt(4)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.accent_color
        line.line.fill.background()

    def get_output_path(self) -> Path:
        """Get the output file path."""
        return self.output_path

"""
Visual Anchor & Script Module
=============================

Converts slides to images, generates scripts,
and extracts mouse coordinates using Vision AI.
"""

import json
import subprocess
import os
from pathlib import Path
from typing import Optional, Dict, Any, List

from rich.console import Console
from rich.panel import Panel
from rich.progress import track, Progress

from ..providers.base import LLMProvider
from ..utils.resumable import check_and_resume, save_json, load_json

console = Console()


class VisualAnchorModule:
    """
    Visual Anchor Module for generating scripts and coordinates.

    This module:
    1. Converts PPTX to PNG images
    2. Uses Vision AI to generate detailed scripts
    3. Extracts mouse cursor coordinates for key moments
    """

    OUTPUT_FILE = "4_script_with_coords.json"
    SLIDES_DIR = "slides_png"

    SCRIPT_SYSTEM_PROMPT = """你是一位專業的教學講師，正在準備課程逐字稿。

你的語調特點：
- 台灣口語化，親切自然
- 適時使用「欸」「對」「那」等語助詞
- 解釋清楚，節奏適中
- 會用「大家可以看到」「這邊我們來看看」等引導語句

請用繁體中文撰寫逐字稿。"""

    SCRIPT_PROMPT_TEMPLATE = """請根據這張投影片圖片，撰寫一段教學逐字稿。

投影片資訊：
- 第 {slide_number} 張投影片
- 章節：{section_title}
- 主題：{topic}

原本的講解重點：
{content}

請撰寫約 100-200 字的逐字稿，語調要：
1. 自然口語化（像在對學生講話）
2. 引導學生注意投影片上的重點
3. 適時停頓讓學生思考

請直接輸出逐字稿內容，不要加標題或其他說明。"""

    COORDS_SYSTEM_PROMPT = """你是一位視覺分析專家，負責分析教學影片中滑鼠游標應該指向的位置。

你的任務是：
1. 閱讀逐字稿
2. 觀察投影片圖片
3. 找出 3-5 個關鍵時刻，滑鼠應該指向的位置

座標格式：
- 使用百分比 (0-100)
- x=0 是最左邊，x=100 是最右邊
- y=0 是最上面，y=100 是最下面"""

    COORDS_PROMPT_TEMPLATE = """請分析這張投影片和對應的逐字稿，找出滑鼠游標應該指向的關鍵位置。

逐字稿：
{script}

請輸出 JSON 格式，包含 3-5 個關鍵點：

```json
{{
  "anchor_points": [
    {{
      "id": 1,
      "trigger_text": "當講到這段話時",
      "description": "指向標題",
      "x_percent": 50,
      "y_percent": 15,
      "duration_ratio": 0.2
    }},
    {{
      "id": 2,
      "trigger_text": "對應的觸發文字",
      "description": "指向的元素描述",
      "x_percent": 數字,
      "y_percent": 數字,
      "duration_ratio": 這個點停留的時間比例
    }}
  ]
}}
```

注意：
1. duration_ratio 所有點加起來應該約等於 1.0
2. 座標要精確指向投影片上的文字或圖表
3. trigger_text 要對應逐字稿中的關鍵句子

請直接輸出 JSON，不要加其他說明。"""

    def __init__(
        self,
        llm_provider: LLMProvider,
        workspace: Path,
        syllabus: Dict[str, Any],
    ):
        """
        Initialize the Visual Anchor Module.

        Args:
            llm_provider: LLM provider for vision analysis
            workspace: Workspace directory
            syllabus: Course syllabus for context
        """
        self.llm = llm_provider
        self.workspace = workspace
        self.syllabus = syllabus
        self.output_path = workspace / self.OUTPUT_FILE
        self.slides_dir = workspace / self.SLIDES_DIR

    def run(
        self,
        pptx_path: Path,
        skip_if_exists: bool = True,
    ) -> Dict[str, Any]:
        """
        Generate scripts and coordinates from slides.

        Args:
            pptx_path: Path to PPTX file
            skip_if_exists: Whether to skip if output exists

        Returns:
            Script with coordinates as dictionary
        """
        console.print(Panel(
            "[bold blue]Step 4: Visual Anchor & Script Generation[/bold blue]\n"
            "Converting slides and generating narration...",
            title="🎯 Visual Anchor Module",
        ))

        # Check for existing output
        if skip_if_exists:
            should_skip, existing_data = check_and_resume(
                self.output_path,
                "Visual Anchor & Script",
                ask_user=True,
                load_func=load_json,
            )
            if should_skip and existing_data:
                return existing_data

        # Step 1: Convert PPTX to images
        image_paths = self._convert_pptx_to_images(pptx_path)

        if not image_paths:
            console.print("[red]Error: No images generated from PPTX[/red]")
            return {"slides": [], "error": "No images generated"}

        # Step 2: Generate scripts and coordinates for each slide
        script_data = {
            "course_title": self.syllabus.get("course_title", "Course"),
            "total_slides": len(image_paths),
            "slides": [],
        }

        # Map slides to sections/topics from syllabus
        slide_contexts = self._build_slide_contexts()

        with Progress() as progress:
            task = progress.add_task(
                "[green]Processing slides...",
                total=len(image_paths)
            )

            for i, image_path in enumerate(image_paths):
                slide_num = i + 1
                context = slide_contexts.get(i, {})

                console.print(f"\n[cyan]Processing slide {slide_num}/{len(image_paths)}...[/cyan]")

                # Generate script
                script = self._generate_script(
                    image_path,
                    slide_num,
                    context.get("section_title", ""),
                    context.get("topic", f"Slide {slide_num}"),
                    context.get("content", ""),
                )

                # Extract coordinates
                coords = self._extract_coordinates(image_path, script)

                slide_data = {
                    "slide_number": slide_num,
                    "image_path": str(image_path),
                    "section_title": context.get("section_title", ""),
                    "topic": context.get("topic", ""),
                    "script": script,
                    "anchor_points": coords.get("anchor_points", []),
                }

                script_data["slides"].append(slide_data)
                progress.update(task, advance=1)

        # Save output
        save_json(script_data, self.output_path)

        console.print("\n[green]✓ Visual anchor and script generation complete![/green]")
        console.print(f"  Output saved to: [cyan]{self.output_path}[/cyan]")
        console.print(f"  Total slides processed: {len(image_paths)}")

        return script_data

    def _convert_pptx_to_images(self, pptx_path: Path) -> List[Path]:
        """Convert PPTX to PNG images using LibreOffice."""
        console.print("[yellow]Converting PPTX to images...[/yellow]")

        # Ensure output directory exists
        self.slides_dir.mkdir(parents=True, exist_ok=True)

        # Check if images already exist
        existing_images = sorted(self.slides_dir.glob("*.png"))
        if existing_images:
            console.print(f"  [dim]Found {len(existing_images)} existing images[/dim]")
            return existing_images

        # Try LibreOffice conversion
        try:
            # First convert to PDF
            pdf_path = self.workspace / "temp_slides.pdf"

            result = subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(self.workspace),
                    str(pptx_path),
                ],
                capture_output=True,
                text=True,
                timeout=120,
            )

            if result.returncode != 0:
                console.print(f"[yellow]LibreOffice conversion warning: {result.stderr}[/yellow]")

            # Find the generated PDF
            pptx_name = pptx_path.stem
            pdf_path = self.workspace / f"{pptx_name}.pdf"

            if pdf_path.exists():
                # Convert PDF to images using pdf2image
                try:
                    from pdf2image import convert_from_path

                    images = convert_from_path(pdf_path, dpi=150)

                    image_paths = []
                    for i, img in enumerate(images):
                        img_path = self.slides_dir / f"page_{i+1:03d}.png"
                        img.save(img_path, "PNG")
                        image_paths.append(img_path)

                    console.print(f"  [green]Generated {len(image_paths)} images[/green]")
                    return image_paths

                except ImportError:
                    console.print("[yellow]pdf2image not available, trying alternative method...[/yellow]")

        except FileNotFoundError:
            console.print("[yellow]LibreOffice not found, trying alternative method...[/yellow]")
        except subprocess.TimeoutExpired:
            console.print("[yellow]LibreOffice conversion timed out[/yellow]")

        # Fallback: Create placeholder images
        return self._create_placeholder_images(pptx_path)

    def _create_placeholder_images(self, pptx_path: Path) -> List[Path]:
        """Create placeholder images if conversion fails."""
        from pptx import Presentation
        from PIL import Image, ImageDraw, ImageFont

        console.print("[yellow]Creating placeholder images from PPTX...[/yellow]")

        prs = Presentation(str(pptx_path))
        image_paths = []

        for i, slide in enumerate(prs.slides):
            img = Image.new("RGB", (1920, 1080), color=(255, 255, 255))
            draw = ImageDraw.Draw(img)

            # Draw slide number
            draw.text(
                (50, 50),
                f"Slide {i+1}",
                fill=(43, 87, 154),
            )

            # Try to extract text from slide
            y_pos = 150
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text = shape.text[:100]  # Limit length
                    draw.text(
                        (50, y_pos),
                        text,
                        fill=(50, 50, 50),
                    )
                    y_pos += 50
                    if y_pos > 900:
                        break

            img_path = self.slides_dir / f"page_{i+1:03d}.png"
            img.save(img_path)
            image_paths.append(img_path)

        console.print(f"  [green]Created {len(image_paths)} placeholder images[/green]")
        return image_paths

    def _build_slide_contexts(self) -> Dict[int, Dict[str, str]]:
        """Build context mapping for each slide from syllabus."""
        contexts = {}
        slide_index = 0

        # Title slide
        contexts[slide_index] = {
            "section_title": "開場",
            "topic": self.syllabus.get("course_title", "課程"),
            "content": self.syllabus.get("course_description", ""),
        }
        slide_index += 1

        # Overview slide
        contexts[slide_index] = {
            "section_title": "課程大綱",
            "topic": "今天我們要學什麼",
            "content": "課程大綱總覽",
        }
        slide_index += 1

        # Section slides
        for section in self.syllabus.get("sections", []):
            # Section intro
            contexts[slide_index] = {
                "section_title": section.get("title", ""),
                "topic": f"Chapter {section.get('section_id', '')}",
                "content": section.get("description", ""),
            }
            slide_index += 1

            # Metaphor slide (if exists)
            if section.get("metaphor"):
                contexts[slide_index] = {
                    "section_title": section.get("title", ""),
                    "topic": "比喻說明",
                    "content": section.get("metaphor", ""),
                }
                slide_index += 1

            # Talking points
            for point in section.get("talking_points", []):
                contexts[slide_index] = {
                    "section_title": section.get("title", ""),
                    "topic": point.get("topic", ""),
                    "content": point.get("content", ""),
                }
                slide_index += 1

        # Summary slide
        contexts[slide_index] = {
            "section_title": "總結",
            "topic": "課程總結",
            "content": self.syllabus.get("summary", ""),
        }

        return contexts

    def _generate_script(
        self,
        image_path: Path,
        slide_num: int,
        section_title: str,
        topic: str,
        content: str,
    ) -> str:
        """Generate narration script for a slide."""
        prompt = self.SCRIPT_PROMPT_TEMPLATE.format(
            slide_number=slide_num,
            section_title=section_title,
            topic=topic,
            content=content or "（請根據投影片內容生成逐字稿）",
        )

        try:
            response = self.llm.vision_analyze(
                image_path,
                prompt,
                temperature=0.7,
            )
            return response.content.strip()
        except Exception as e:
            console.print(f"[red]Warning: Script generation failed: {e}[/red]")
            return f"這是第 {slide_num} 張投影片，主題是 {topic}。{content}"

    def _extract_coordinates(
        self,
        image_path: Path,
        script: str,
    ) -> Dict[str, Any]:
        """Extract mouse coordinates for key moments."""
        prompt = self.COORDS_PROMPT_TEMPLATE.format(script=script)

        try:
            response = self.llm.vision_analyze(
                image_path,
                prompt,
                temperature=0.3,  # Lower temp for more consistent output
            )

            # Parse JSON response
            content = response.content

            # Extract JSON from response
            if "```json" in content:
                start = content.find("```json") + 7
                end = content.find("```", start)
                if end > start:
                    content = content[start:end]
            elif "```" in content:
                start = content.find("```") + 3
                end = content.find("```", start)
                if end > start:
                    content = content[start:end]

            return json.loads(content.strip())

        except Exception as e:
            console.print(f"[yellow]Warning: Coordinate extraction failed: {e}[/yellow]")
            # Return default coordinates
            return {
                "anchor_points": [
                    {
                        "id": 1,
                        "trigger_text": "開始",
                        "description": "標題位置",
                        "x_percent": 50,
                        "y_percent": 20,
                        "duration_ratio": 0.3,
                    },
                    {
                        "id": 2,
                        "trigger_text": "內容",
                        "description": "主要內容",
                        "x_percent": 50,
                        "y_percent": 50,
                        "duration_ratio": 0.5,
                    },
                    {
                        "id": 3,
                        "trigger_text": "結束",
                        "description": "底部",
                        "x_percent": 50,
                        "y_percent": 80,
                        "duration_ratio": 0.2,
                    },
                ]
            }

    def get_output_path(self) -> Path:
        """Get the output file path."""
        return self.output_path

    def get_slides_dir(self) -> Path:
        """Get the slides image directory."""
        return self.slides_dir
